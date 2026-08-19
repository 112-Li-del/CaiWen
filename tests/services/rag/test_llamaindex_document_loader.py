"""Tests for LlamaIndex document loading."""

from __future__ import annotations

import asyncio
import base64
import hashlib
import io
import json
from pathlib import Path
from types import SimpleNamespace

from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
import pymupdf
import pytest


def _make_docx(paragraphs: list[str]) -> bytes:
    doc = DocxDocument()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _make_xlsx(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _make_pptx(texts: list[str]) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    for index, text in enumerate(texts):
        text_box = slide.shapes.add_textbox(
            Inches(1),
            Inches(1 + index * 0.5),
            Inches(6),
            Inches(0.5),
        )
        text_box.text_frame.text = text
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _make_pdf(path: Path, page_texts: list[str]) -> None:
    with pymupdf.open() as pdf:
        for text in page_texts:
            page = pdf.new_page()
            page.insert_text((72, 72), text)
        pdf.save(path)


_ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _make_pdf_with_images(path: Path, page_texts: list[str | None]) -> None:
    with pymupdf.open() as pdf:
        for text in page_texts:
            page = pdf.new_page()
            page.insert_image(pymupdf.Rect(72, 100, 144, 172), stream=_ONE_PIXEL_PNG)
            if text:
                page.insert_text((72, 72), text)
        pdf.save(path)


def test_loader_extracts_chat_supported_office_files(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    pytest.importorskip("llama_index.core")
    from deeptutor.knowledge.renditions import Rendition, RenditionManager
    from deeptutor.services.rag.pipelines.llamaindex.document_loader import (
        LlamaIndexDocumentLoader,
    )

    docx_path = tmp_path / "notes.docx"
    docx_path.write_bytes(_make_docx(["Docx paragraph"]))
    xlsx_path = tmp_path / "book.xlsx"
    xlsx_path.write_bytes(_make_xlsx([["cell-a", 42]]))
    pptx_path = tmp_path / "slides.pptx"
    pptx_path.write_bytes(_make_pptx(["Slide title", "Slide body"]))
    preview_path = tmp_path / "docx-preview.pdf"
    _make_pdf(preview_path, ["Docx page one", "Docx page two"])

    def _fake_ensure_pdf(_self, original_path: Path) -> Rendition:
        digest = hashlib.sha256(original_path.read_bytes()).hexdigest()
        return Rendition(
            original_path=original_path,
            pdf_path=preview_path,
            preview_path=f"renditions/{digest}.pdf",
            original_sha256=digest,
            preview_sha256="preview-hash",
            page_count=2,
            converted_from_docx=True,
        )

    monkeypatch.setattr(RenditionManager, "ensure_pdf", _fake_ensure_pdf)

    documents = asyncio.run(
        LlamaIndexDocumentLoader().load([str(docx_path), str(xlsx_path), str(pptx_path)])
    )

    assert {doc.metadata["file_name"] for doc in documents} == {
        "notes.docx",
        "book.xlsx",
        "slides.pptx",
    }
    all_text = "\n".join(doc.text for doc in documents)
    assert "Docx page one" in all_text
    assert "Docx page two" in all_text
    assert "cell-a" in all_text
    assert "42" in all_text
    assert "Slide title" in all_text
    assert "Slide body" in all_text

    docx_documents = [doc for doc in documents if doc.metadata["file_name"] == "notes.docx"]
    assert [doc.metadata["page"] for doc in docx_documents] == [1, 2]
    assert all(doc.metadata["page_count"] == 2 for doc in docx_documents)
    assert all(doc.metadata["preview_mime"] == "application/pdf" for doc in docx_documents)
    assert all(doc.metadata["converted_from_docx"] is True for doc in docx_documents)


def test_retrieval_source_exposes_page_and_relative_paths_only() -> None:
    pytest.importorskip("llama_index.core")
    from deeptutor.services.rag.pipelines.llamaindex.pipeline import LlamaIndexPipeline

    pipeline = LlamaIndexPipeline.__new__(LlamaIndexPipeline)
    node = SimpleNamespace(
        score=0.95,
        node=SimpleNamespace(
            text="The cited school rule.",
            node_id="chunk-7",
            metadata={
                "file_name": "student-handbook.docx",
                "file_path": r"C:\\secret\\school\\student-handbook.docx",
                "raw_name": "student-handbook.docx",
                "preview_path": "renditions/abc.pdf",
                "preview_mime": "application/pdf",
                "converted_from_docx": True,
                "page_number": 7,
                "page_count": 42,
            },
        ),
    )

    result = pipeline._nodes_to_result("rule", [node])
    source = result["sources"][0]

    assert source["source_file"] == "student-handbook.docx"
    assert source["preview_path"] == "renditions/abc.pdf"
    assert source["page"] == 7
    assert source["page_count"] == 42
    assert source["chunk_id"] == "chunk-7"
    assert "C:\\secret" not in json.dumps(source)


def test_loader_rejects_image_only_pdf_and_records_document_status(tmp_path: Path) -> None:
    pytest.importorskip("llama_index.core")
    from deeptutor.knowledge.renditions import RenditionManager
    from deeptutor.services.rag.pipelines.llamaindex.document_loader import (
        ImageOnlyDocumentError,
        LlamaIndexDocumentLoader,
    )

    source = tmp_path / "scan.pdf"
    _make_pdf_with_images(source, [None, None])

    loader = LlamaIndexDocumentLoader()
    with pytest.raises(ImageOnlyDocumentError, match="文档全是图片，无法建立索引"):
        asyncio.run(loader.load([str(source)]))

    inspection = RenditionManager(tmp_path).entry_for(source.name)
    assert inspection is not None
    assert inspection["indexing_status"] == "image_only"
    assert inspection["text_page_count"] == 0
    assert inspection["image_page_count"] == 2
    assert loader.last_load_summary["total_file_count"] == 1
    assert loader.last_load_summary["indexed_file_count"] == 0
    assert loader.last_load_summary["image_only_file_count"] == 1
    assert loader.last_load_summary["skipped_file_count"] == 1


def test_loader_indexes_text_in_mixed_pdf_and_records_warning_status(tmp_path: Path) -> None:
    pytest.importorskip("llama_index.core")
    from deeptutor.knowledge.renditions import RenditionManager
    from deeptutor.services.rag.pipelines.llamaindex.document_loader import (
        LlamaIndexDocumentLoader,
    )

    source = tmp_path / "mixed.pdf"
    _make_pdf_with_images(source, ["School rule", None])

    loader = LlamaIndexDocumentLoader()
    documents = asyncio.run(loader.load([str(source)]))

    assert len(documents) == 1
    inspection = RenditionManager(tmp_path).entry_for(source.name)
    assert inspection is not None
    assert inspection["indexing_status"] == "partial_images"
    assert inspection["text_page_count"] == 1
    assert inspection["image_page_count"] == 2
    assert loader.last_load_summary["total_file_count"] == 1
    assert loader.last_load_summary["indexed_file_count"] == 1
    assert loader.last_load_summary["image_only_file_count"] == 0
    assert loader.last_load_summary["skipped_file_count"] == 0
